import openpyxl
from typing import List
from apps.testing_analysis.parsers.base_parser import BaseParser
from apps.testing_analysis.schemas import TestCaseRecord

class ExcelParser(BaseParser):
    def parse(self, file_path: str) -> List[TestCaseRecord]:
        records = []
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        
        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            header_row_idx = None
            col_map = {} # mapped field name -> index (0-based)
            
            # Quét tìm header
            for row_idx, row in enumerate(sheet.iter_rows(values_only=True)):
                row_str = " ".join([str(cell).lower() for cell in row if cell])
                # heuristic: if we see some keywords, it's a header
                if "id" in row_str or "status" in row_str or "result" in row_str:
                    header_row_idx = row_idx
                    for col_idx, cell_value in enumerate(row):
                        mapped_name = self._normalize_column_name(cell_value)
                        if mapped_name:
                            col_map[mapped_name] = col_idx
                    break
            
            if header_row_idx is None or not col_map:
                continue # Bỏ qua sheet này nếu không thấy header
                
            # Đọc dữ liệu
            for row_idx, row in enumerate(sheet.iter_rows(values_only=True)):
                if row_idx <= header_row_idx:
                    continue
                
                # Check if row is entirely empty
                if not any(row):
                    continue
                    
                record_data = {}
                for field_name, col_idx in col_map.items():
                    if col_idx < len(row):
                        val = row[col_idx]
                        record_data[field_name] = str(val).strip() if val is not None else None
                        
                # Only add if at least an ID or Description exists
                if record_data.get("test_case_id") or record_data.get("description"):
                    records.append(TestCaseRecord(**record_data))
                    
        wb.close()
        return records
